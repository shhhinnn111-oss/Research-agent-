# -*- coding: utf-8 -*-
"""
Render the WFP Sindh briefing note to PDF (reportlab) and DOCX (python-docx).
Usage: python3 build_brief.py
Outputs: WFP_Sindh_Brief.pdf, WFP_Sindh_Brief.docx
"""
import re
import sys

from brief_content import TITLE, SUBTITLE, META, BLOCKS

# ----------------------------------------------------------------------------- PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                Table, TableStyle, KeepTogether)

FONT = "Helvetica"
FONT_B = "Helvetica-Bold"
FONT_I = "Helvetica-Oblique"
FONT_BI = "Helvetica-BoldOblique"

NAVY = colors.HexColor("#0B3A5D")
GREY = colors.HexColor("#555555")
RULE = colors.HexColor("#0B3A5D")
ROW_ALT = colors.HexColor("#F1F5F9")

BASE = 8.6          # body font size
LEAD = 10.3         # leading


def styles():
    s = {}
    s["title"] = ParagraphStyle("title", fontName=FONT_B, fontSize=12.5, leading=14.5,
                                textColor=NAVY, spaceAfter=1)
    s["subtitle"] = ParagraphStyle("subtitle", fontName=FONT_I, fontSize=9, leading=11,
                                   textColor=GREY, spaceAfter=0.5)
    s["meta"] = ParagraphStyle("meta", fontName=FONT, fontSize=7.2, leading=8.6,
                               textColor=GREY, spaceAfter=3)
    s["h"] = ParagraphStyle("h", fontName=FONT_B, fontSize=9.4, leading=11.2, textColor=NAVY,
                            spaceBefore=3.6, spaceAfter=1.4, keepWithNext=1)
    s["p"] = ParagraphStyle("p", fontName=FONT, fontSize=BASE, leading=LEAD, alignment=TA_JUSTIFY,
                            spaceAfter=2.2)
    s["b"] = ParagraphStyle("b", parent=s["p"], leftIndent=8, bulletIndent=1, spaceAfter=1.4,
                            bulletFontName=FONT, bulletFontSize=BASE)
    s["cell"] = ParagraphStyle("cell", fontName=FONT, fontSize=7.5, leading=8.6, alignment=TA_LEFT)
    s["cellb"] = ParagraphStyle("cellb", parent=s["cell"], fontName=FONT_B)
    s["cellh"] = ParagraphStyle("cellh", parent=s["cell"], fontName=FONT_B, textColor=colors.white)
    s["footer"] = ParagraphStyle("footer", fontName=FONT, fontSize=6.8, leading=8, textColor=GREY)
    return s


def esc(t):
    return t.replace("&", "&amp;")


def build_pdf(path):
    s = styles()
    doc = BaseDocTemplate(path, pagesize=A4,
                          leftMargin=14 * mm, rightMargin=14 * mm,
                          topMargin=10 * mm, bottomMargin=10.5 * mm,
                          title=TITLE, author="WFP Pakistan Country Office", subject=SUBTITLE)
    W, H = A4
    frame = Frame(doc.leftMargin, doc.bottomMargin, W - doc.leftMargin - doc.rightMargin,
                  H - doc.topMargin - doc.bottomMargin, id="f", leftPadding=0, rightPadding=0,
                  topPadding=0, bottomPadding=0)

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setFont(FONT, 6.8)
        canvas.setFillColor(GREY)
        canvas.drawString(doc.leftMargin, 6.5 * mm,
                          "WFP Pakistan  |  Sindh briefing note  |  Internal - not for onward distribution")
        canvas.drawRightString(W - doc.rightMargin, 6.5 * mm, "Page %d of 2" % d.page)
        canvas.setStrokeColor(RULE)
        canvas.setLineWidth(0.4)
        canvas.line(doc.leftMargin, 9 * mm, W - doc.rightMargin, 9 * mm)
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=on_page)])

    story = [Paragraph(esc(TITLE), s["title"]), Paragraph(esc(SUBTITLE), s["subtitle"]), Paragraph(esc(META), s["meta"])]
    # rule under header
    story.append(Table([[""]], colWidths=[W - doc.leftMargin - doc.rightMargin], rowHeights=[1.2],
                       style=TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.8, RULE)])))
    story.append(Spacer(1, 1.5))

    for blk in BLOCKS:
        kind = blk[0]
        if kind == "h":
            story.append(Paragraph(esc(blk[1]), s["h"]))
        elif kind == "p":
            story.append(Paragraph(esc(blk[1]), s["p"]))
        elif kind == "bullets":
            for item in blk[1]:
                story.append(Paragraph(esc(item), s["b"], bulletText="\u2022"))
        elif kind == "table":
            header, rows, fracs = blk[1], blk[2], blk[3]
            avail = W - doc.leftMargin - doc.rightMargin
            colw = [f * avail for f in fracs]
            data = [[Paragraph(esc(h), s["cellh"]) for h in header]]
            for r in rows:
                data.append([Paragraph(esc(r[0]), s["cellb"]), Paragraph(esc(r[1]), s["cell"]), Paragraph(esc(r[2]), s["cell"])])
            t = Table(data, colWidths=colw, repeatRows=1)
            ts = [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 1.2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D3DD")),
                ("BOX", (0, 0), (-1, -1), 0.5, NAVY),
            ]
            for i in range(1, len(data)):
                if i % 2 == 0:
                    ts.append(("BACKGROUND", (0, i), (-1, i), ROW_ALT))
            t.setStyle(TableStyle(ts))
            story.append(t)
            story.append(Spacer(1, 2))
        elif kind == "footer":
            story.append(Paragraph(esc(blk[1]), s["footer"]))
    doc.build(story)


# ----------------------------------------------------------------------------- DOCX
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TAG_RE = re.compile(r"(<b>|</b>|<i>|</i>)")


def add_rich(par, text, size, bold_base=False, color=None):
    bold = bold_base
    ital = False
    for tok in TAG_RE.split(text):
        if tok == "<b>":
            bold = True
        elif tok == "</b>":
            bold = bold_base
        elif tok == "<i>":
            ital = True
        elif tok == "</i>":
            ital = False
        elif tok:
            run = par.add_run(tok)
            run.bold = bold
            run.italic = ital
            run.font.size = Pt(size)
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
            if color is not None:
                run.font.color.rgb = color


def shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(table, top=20, bottom=20, left=50, right=50):
    tblPr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:%s" % k)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def para_fmt(par, before=0, after=2, line=1.0, align=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        par.alignment = align


def build_docx(path):
    navy = RGBColor(0x0B, 0x3A, 0x5D)
    grey = RGBColor(0x55, 0x55, 0x55)
    body = 8.5
    d = Document()
    sec = d.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin = Cm(1.1)
    sec.bottom_margin = Cm(1.1)
    sec.footer_distance = Cm(0.5)

    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(body)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    p = d.add_paragraph()
    add_rich(p, TITLE, 12.5, bold_base=True, color=navy)
    para_fmt(p, after=0)
    p = d.add_paragraph()
    add_rich(p, "<i>%s</i>" % SUBTITLE, 9, color=grey)
    para_fmt(p, after=0)
    p = d.add_paragraph()
    add_rich(p, META, 7.2, color=grey)
    para_fmt(p, after=2)
    # bottom border on meta paragraph
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0B3A5D")
    pbdr.append(bottom)
    pPr.append(pbdr)

    for blk in BLOCKS:
        kind = blk[0]
        if kind == "h":
            p = d.add_paragraph()
            add_rich(p, blk[1], 9.4, bold_base=True, color=navy)
            para_fmt(p, before=3.5, after=1.2)
            p.paragraph_format.keep_with_next = True
        elif kind == "p":
            p = d.add_paragraph()
            add_rich(p, blk[1], body)
            para_fmt(p, after=2.2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "bullets":
            for item in blk[1]:
                p = d.add_paragraph(style="List Bullet")
                add_rich(p, item, body)
                para_fmt(p, after=1.2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                p.paragraph_format.left_indent = Cm(0.45)
                p.paragraph_format.first_line_indent = Cm(-0.3)
        elif kind == "table":
            header, rows, fracs = blk[1], blk[2], blk[3]
            avail_cm = 21.0 - 2 * 1.4
            t = d.add_table(rows=1, cols=3)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            set_cell_margins(t)
            widths = [Cm(f * avail_cm) for f in fracs]
            for i, h in enumerate(header):
                c = t.rows[0].cells[i]
                c.width = widths[i]
                c.paragraphs[0].text = ""
                add_rich(c.paragraphs[0], h, 7.5, bold_base=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                para_fmt(c.paragraphs[0], after=0)
                shade(c, "0B3A5D")
            for ri, r in enumerate(rows):
                cells = t.add_row().cells
                for i, txt in enumerate(r):
                    cells[i].width = widths[i]
                    cells[i].paragraphs[0].text = ""
                    add_rich(cells[i].paragraphs[0], txt, 7.5, bold_base=(i == 0))
                    para_fmt(cells[i].paragraphs[0], after=0)
                    if ri % 2 == 1:
                        shade(cells[i], "F1F5F9")
            # spacer
            p = d.add_paragraph()
            para_fmt(p, after=0)
            for r in p.runs:
                r.font.size = Pt(2)
            p.paragraph_format.line_spacing = Pt(3)
        elif kind == "footer":
            p = d.add_paragraph()
            add_rich(p, blk[1], 6.8, color=grey)
            para_fmt(p, after=0)

    # footer text
    fp = sec.footer.paragraphs[0]
    add_rich(fp, "WFP Pakistan  |  Sindh briefing note  |  Internal - not for onward distribution", 6.8, color=grey)
    para_fmt(fp, after=0)
    d.save(path)


if __name__ == "__main__":
    out_pdf = "WFP_Sindh_Brief.pdf"
    out_docx = "WFP_Sindh_Brief.docx"
    build_pdf(out_pdf)
    build_docx(out_docx)
    import pymupdf
    n = len(pymupdf.open(out_pdf))
    print("PDF pages:", n)
    sys.exit(0 if n == 2 else 1)
